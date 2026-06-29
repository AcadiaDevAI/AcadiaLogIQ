"""Unit tests for the dedicated Excel parser (backend.ingestion.xlsx_parser).

Exercises the lossless-ingestion guarantees on real in-memory .xlsx files:
header auto-detection past a title row, full column width, multiple sheets,
number/date stringification, and skipping only truly-empty rows. Only .xlsx
is exercised (openpyxl is a hard dep); .xls would also require xlrd.

Run with: py -m unittest backend.tests.test_excel_parser
"""
from __future__ import annotations

import datetime
import os
import tempfile
import unittest
from pathlib import Path

os.environ.setdefault("DATABASE_URL", "postgresql+psycopg://u:p@localhost:5432/test")
os.environ.setdefault("AWS_SECRETS_DISABLED", "true")

from openpyxl import Workbook  # noqa: E402

from backend.ingestion.xlsx_parser import parse_excel  # noqa: E402


class TestXlsxParser(unittest.TestCase):
    def setUp(self):
        self._paths = []

    def tearDown(self):
        for p in self._paths:
            try:
                p.unlink()
            except OSError:
                pass

    def _save(self, wb) -> Path:
        fd, name = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)
        wb.save(name)
        path = Path(name)
        self._paths.append(path)
        return path

    def _simple(self, rows):
        wb = Workbook()
        ws = wb.active
        for r in rows:
            ws.append(r)
        return self._save(wb)

    def test_title_row_then_header_multi_column(self):
        path = self._simple([
            ["Walgreens LogIQ Demo Review"],
            [],
            ["Ticket_ID", "Customer", "Priority", "Summary"],
            ["INC-1", "Acme", "P1", "Login fails"],
            ["INC-2", "Globex", "P2", "Slow API"],
        ])
        chunks = parse_excel(path)
        rows = [c for c in chunks if c.chunk_type == "xlsx_row"]
        self.assertEqual(len(rows), 2)
        c0 = rows[0]
        for frag in ("Ticket_ID: INC-1", "Customer: Acme",
                     "Priority: P1", "Summary: Login fails"):
            self.assertIn(frag, c0.text)
        self.assertEqual(c0.metadata.get("primary_id"), "INC-1")
        self.assertEqual(c0.metadata.get("customer"), "Acme")
        self.assertEqual(c0.section_heading, "INC-1")
        ctx = [c for c in chunks if c.chunk_type == "xlsx_context"]
        self.assertTrue(any("Walgreens LogIQ Demo Review" in c.text for c in ctx))

    def test_multiple_sheets(self):
        wb = Workbook()
        ws1 = wb.active
        ws1.title = "Tickets"
        ws1.append(["id", "note"])
        ws1.append(["1", "alpha"])
        ws2 = wb.create_sheet("Contacts")
        ws2.append(["name", "email"])
        ws2.append(["Bob", "bob@x.com"])
        path = self._save(wb)
        chunks = parse_excel(path)
        rows = [c for c in chunks if c.chunk_type == "xlsx_row"]
        self.assertEqual(len(rows), 2)
        self.assertEqual({c.metadata.get("sheet") for c in rows}, {"Tickets", "Contacts"})

    def test_numbers_and_dates_stringified(self):
        path = self._simple([
            ["id", "count", "opened"],
            ["A", 42, datetime.datetime(2026, 1, 15, 9, 30)],
        ])
        row = [c for c in parse_excel(path) if c.chunk_type == "xlsx_row"][0]
        self.assertIn("count: 42", row.text)
        self.assertIn("2026-01-15", row.text)

    def test_ragged_row_extra_cell_kept(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["a", "b"])
        ws.append(["x", "y", "EXTRA"])
        path = self._save(wb)
        row = [c for c in parse_excel(path) if c.chunk_type == "xlsx_row"][0]
        self.assertIn("EXTRA", row.text)

    def test_blank_rows_skipped(self):
        path = self._simple([
            ["id", "note"],
            ["1", "real"],
            ["", ""],
            [None, None],
        ])
        rows = [c for c in parse_excel(path) if c.chunk_type == "xlsx_row"]
        self.assertEqual(len(rows), 1)
        self.assertIn("note: real", rows[0].text)

    def test_empty_workbook(self):
        wb = Workbook()
        path = self._save(wb)
        self.assertEqual(parse_excel(path), [])


if __name__ == "__main__":
    unittest.main()
